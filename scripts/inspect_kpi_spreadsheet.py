import openpyxl
import json

wb = openpyxl.load_workbook('kpi_reestr_1vzWZzdk.xlsx', data_only=True)

report = {
    "total_sheets": len(wb.sheetnames),
    "sheets": []
}

for sheet_name in wb.sheetnames:
    ws = wb[sheet_name]
    headers = []
    sample_rows = []
    
    # find first non-empty header row
    header_row_idx = 1
    for r in range(1, min(6, ws.max_row + 1)):
        row_vals = [str(ws.cell(r, c).value or '').strip() for c in range(1, ws.max_column + 1)]
        non_empty = [v for v in row_vals if v]
        if len(non_empty) >= 3:
            header_row_idx = r
            headers = row_vals[:20]
            break
            
    # get sample rows
    for r in range(header_row_idx + 1, min(header_row_idx + 6, ws.max_row + 1)):
        row_vals = [str(ws.cell(r, c).value or '').strip() for c in range(1, len(headers) + 1)]
        if any(row_vals):
            sample_rows.append(row_vals)
            
    report["sheets"].append({
        "name": sheet_name,
        "max_rows": ws.max_row,
        "max_cols": ws.max_column,
        "header_row": header_row_idx,
        "headers": [h for h in headers if h],
        "sample_rows_count": len(sample_rows),
        "sample_first_row": sample_rows[0] if sample_rows else []
    })

with open('kpi_spreadsheet_analysis.json', 'w', encoding='utf-8') as f:
    json.dump(report, f, ensure_ascii=False, indent=2)

print("Analysis saved to kpi_spreadsheet_analysis.json")
